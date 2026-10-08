import pandas as pd
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter

# ============================================================
# PATHS
# ============================================================

OUTPUT_DIR = Path("excel")
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

OUTPUT_FILE = OUTPUT_DIR / "FinSight_360_TCS_Financial_Model.xlsx"

RESULTS = Path("results/tables")

# ============================================================
# LOAD VALIDATED PROJECT OUTPUTS
# ============================================================

historical = pd.read_csv(
    RESULTS / "tcs_historical_financial_model_base.csv"
)

forecast = pd.read_csv(
    RESULTS / "tcs_scenario_forecast_FY2027_FY2031.csv"
)

dcf = pd.read_csv(
    RESULTS / "tcs_dcf_final_equity_valuation.csv"
)

sensitivity = pd.read_csv(
    RESULTS / "tcs_dcf_per_share_sensitivity.csv",
    index_col=0
)

relative = pd.read_csv(
    RESULTS / "fy2026_peer_relative_valuation.csv"
)

fpa = pd.read_csv(
    RESULTS / "tcs_fpa_budget_variance.csv"
)

treasury = pd.read_csv(
    RESULTS / "tcs_treasury_liquidity_kpis.csv"
)

credit = pd.read_csv(
    RESULTS / "tcs_credit_scorecard.csv"
)

# ============================================================
# WORKBOOK
# ============================================================

wb = Workbook()

default = wb.active
wb.remove(default)

sheet_names = [
    "Cover",
    "Historical",
    "Forecast",
    "DCF",
    "Sensitivity",
    "Comps",
    "FP&A",
    "Treasury",
    "Credit"
]

for name in sheet_names:
    wb.create_sheet(name)

# ============================================================
# STYLING
# ============================================================

dark_fill = PatternFill(
    "solid",
    fgColor="1F4E78"
)

section_fill = PatternFill(
    "solid",
    fgColor="D9EAF7"
)

input_fill = PatternFill(
    "solid",
    fgColor="FFF2CC"
)

thin = Side(
    style="thin",
    color="B7B7B7"
)

border = Border(
    left=thin,
    right=thin,
    top=thin,
    bottom=thin
)

def title(ws, text):
    ws.merge_cells("A1:H2")
    c = ws["A1"]
    c.value = text
    c.font = Font(
        size=18,
        bold=True,
        color="FFFFFF"
    )
    c.fill = dark_fill
    c.alignment = Alignment(
        horizontal="center",
        vertical="center"
    )

def format_table(ws):
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None:
                cell.border = border
                cell.alignment = Alignment(
                    vertical="center"
                )

    for cell in ws[4]:
        if cell.value is not None:
            cell.font = Font(
                bold=True,
                color="FFFFFF"
            )
            cell.fill = dark_fill

    for col in range(
        1,
        ws.max_column + 1
    ):
        width = 12

        for cell in ws[
            get_column_letter(col)
        ]:
            if cell.value is not None:
                width = max(
                    width,
                    min(
                        len(str(cell.value)) + 2,
                        28
                    )
                )

        ws.column_dimensions[
            get_column_letter(col)
        ].width = width

    ws.freeze_panes = "A5"

def write_df(ws, df, start_row=4):

    for col_num, col in enumerate(
        df.columns,
        1
    ):
        ws.cell(
            start_row,
            col_num,
            col
        )

    for row_num, row in enumerate(
        df.itertuples(
            index=False
        ),
        start_row + 1
    ):
        for col_num, value in enumerate(
            row,
            1
        ):
            ws.cell(
                row_num,
                col_num,
                value
            )

# ============================================================
# COVER
# ============================================================

ws = wb["Cover"]

title(
    ws,
    "FinSight 360 — TCS Financial Analysis & Valuation Model"
)

cover_rows = [
    ("Primary Company", "Tata Consultancy Services"),
    ("Peer Companies", "Infosys, HCL Technologies"),
    ("Historical Period", "FY2022–FY2026"),
    ("Forecast Period", "FY2027–FY2031"),
    ("Currency", "INR Crore unless stated otherwise"),
    ("Valuation Methods", "DCF + Relative Valuation"),
    ("Market Reference Date", "06-Oct-2026"),
    ("Project", "FinSight 360"),
]

r = 4

for label, value in cover_rows:

    ws.cell(r, 1, label)
    ws.cell(r, 2, value)

    ws.cell(r, 1).font = Font(
        bold=True
    )

    ws.cell(r, 1).fill = section_fill

    r += 1

ws["A14"] = "Model Notes"
ws["A14"].font = Font(bold=True)
ws["A14"].fill = section_fill

notes = [
    "Historical financials are based on consolidated annual-report data.",
    "EBITDA is an analytical derived metric where not directly reported.",
    "FY2027–FY2031 figures are scenario-based analytical forecasts, not company guidance.",
    "DCF discount rate and terminal growth rates are model assumptions.",
    "Relative valuation currently uses P/E; unsupported peer EV/EBITDA inputs are not fabricated.",
    "Restricted-purpose financial assets are excluded from TCS available liquidity.",
    "Credit scorecard is an internal analytical framework, not an external credit rating."
]

for i, note in enumerate(
    notes,
    15
):
    ws.cell(i, 1, "• " + note)

ws.column_dimensions["A"].width = 35
ws.column_dimensions["B"].width = 55

# ============================================================
# HISTORICAL
# ============================================================

ws = wb["Historical"]
title(ws, "TCS Historical Financial Model")
write_df(ws, historical)
format_table(ws)

# ============================================================
# FORECAST
# ============================================================

ws = wb["Forecast"]
title(ws, "TCS Scenario Forecast — FY2027 to FY2031")
write_df(ws, forecast)
format_table(ws)

# Highlight assumptions / forecasts
for row in ws.iter_rows(
    min_row=5
):
    for cell in row:
        cell.fill = input_fill

# ============================================================
# DCF
# ============================================================

ws = wb["DCF"]
title(ws, "TCS Discounted Cash Flow Valuation")
write_df(ws, dcf)
format_table(ws)

# ============================================================
# SENSITIVITY
# ============================================================

ws = wb["Sensitivity"]
title(ws, "TCS DCF Per-Share Sensitivity")

sens = sensitivity.reset_index()

write_df(
    ws,
    sens
)

format_table(ws)

# ============================================================
# COMPS
# ============================================================

ws = wb["Comps"]
title(ws, "FY2026 Peer Relative Valuation")
write_df(ws, relative)
format_table(ws)

# ============================================================
# FP&A
# ============================================================

ws = wb["FP&A"]
title(ws, "TCS FP&A — FY2026 Actual vs FY2027 Base Budget")
write_df(ws, fpa)
format_table(ws)

# ============================================================
# TREASURY
# ============================================================

ws = wb["Treasury"]
title(ws, "TCS Treasury & Liquidity Analysis")
write_df(ws, treasury)
format_table(ws)

# ============================================================
# CREDIT
# ============================================================

ws = wb["Credit"]
title(ws, "TCS Internal Corporate Credit Scorecard")
write_df(ws, credit)
format_table(ws)

# ============================================================
# NUMBER FORMATTING
# ============================================================

for ws in wb.worksheets:

    for row in ws.iter_rows():

        for cell in row:

            if isinstance(
                cell.value,
                float
            ):
                cell.number_format = (
                    '#,##0.00'
                )

# ============================================================
# WORKBOOK PROPERTIES
# ============================================================

wb.calculation.fullCalcOnLoad = True
wb.calculation.forceFullCalc = True
wb.calculation.calcMode = "auto"

# ============================================================
# SAVE
# ============================================================

wb.save(
    OUTPUT_FILE
)

print("\n" + "=" * 100)
print("FINSIGHT 360 — EXCEL FINANCIAL MODEL")
print("=" * 100)

print(
    f"\nWorkbook created -> "
    f"{OUTPUT_FILE}"
)

print(
    f"Sheets created   -> "
    f"{len(sheet_names)}"
)

for name in sheet_names:
    print(f"  ✓ {name}")

print(
    "\nHistorical, Forecast, DCF, Sensitivity, "
    "Comps, FP&A, Treasury and Credit outputs "
    "successfully consolidated."
)

print(
    "\nExcel model generation completed successfully."
)
