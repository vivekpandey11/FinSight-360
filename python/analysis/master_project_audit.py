from pathlib import Path
import pandas as pd
import numpy as np

ROOT = Path(".")
R = ROOT / "results" / "tables"
PBI = ROOT / "powerbi" / "data"
EXCEL = ROOT / "excel"

audit = []

def record(area, test, status, detail=""):
    audit.append({
        "area": area,
        "test": test,
        "status": status,
        "detail": str(detail)
    })

def read_csv(name):
    path = R / name

    if not path.exists():
        record("FILES", name, "FAIL", "Missing file")
        return None

    try:
        df = pd.read_csv(path)
        record(
            "FILES", name, "PASS",
            f"{len(df)} rows x {len(df.columns)} cols"
        )
        return df
    except Exception as e:
        record("FILES", name, "FAIL", str(e))
        return None

def find_col(df, keywords, exclude=None):
    exclude = exclude or []

    for col in df.columns:
        name = col.lower()

        if (
            all(k.lower() in name for k in keywords)
            and not any(e.lower() in name for e in exclude)
        ):
            return col

    return None

print("\n" + "=" * 100)
print("FINSIGHT 360 — MASTER PROJECT AUDIT")
print("=" * 100)

# ============================================================
# 1. SOURCE FILES
# ============================================================

pdf_root = ROOT / "data" / "raw" / "financials"

expected_pdfs = {
    "TCS": 3,
    "INFY": 3,
    "HCLTECH": 3
}

for company, expected in expected_pdfs.items():

    folder = pdf_root / company
    count = len(list(folder.glob("*.pdf"))) if folder.exists() else 0

    record(
        "SOURCE",
        f"{company} annual reports",
        "PASS" if count == expected else "FAIL",
        f"Expected {expected}, found {count}"
    )

# ============================================================
# 2. LOAD CORE OUTPUTS
# ============================================================

hist = read_csv(
    "tcs_historical_financial_model_base.csv"
)

forecast = read_csv(
    "tcs_scenario_forecast_FY2027_FY2031.csv"
)

dcf_ev = read_csv(
    "tcs_dcf_enterprise_value.csv"
)

dcf_final = read_csv(
    "tcs_dcf_final_equity_valuation.csv"
)

sens = read_csv(
    "tcs_dcf_sensitivity_long.csv"
)

comps = read_csv(
    "fy2026_peer_relative_valuation.csv"
)

relative = read_csv(
    "tcs_relative_valuation_summary.csv"
)

portfolio = read_csv(
    "portfolio_strategy_summary.csv"
)

allocation = read_csv(
    "portfolio_10lakh_allocation.csv"
)

risk = read_csv(
    "portfolio_risk_metrics.csv"
)

stress = read_csv(
    "portfolio_stress_tests.csv"
)

fpa = read_csv(
    "tcs_fpa_budget_variance.csv"
)

treasury = read_csv(
    "tcs_treasury_liquidity_kpis.csv"
)

treasury_stress = read_csv(
    "tcs_treasury_liquidity_stress.csv"
)

credit = read_csv(
    "tcs_credit_scorecard.csv"
)

credit_summary = read_csv(
    "tcs_credit_risk_summary.csv"
)

# ============================================================
# 3. HISTORICAL MODEL
# ============================================================

if hist is not None:

    record(
        "HISTORICAL",
        "Five historical years",
        "PASS" if len(hist) == 5 else "FAIL",
        f"Rows = {len(hist)}"
    )

    if "revenue" in hist.columns:

        valid = (
            pd.to_numeric(
                hist["revenue"],
                errors="coerce"
            ) > 0
        ).all()

        record(
            "HISTORICAL",
            "Revenue positive",
            "PASS" if valid else "FAIL"
        )

    if {
        "cash_from_operations",
        "capex",
        "fcf"
    }.issubset(hist.columns):

        calc = (
            hist["cash_from_operations"]
            - hist["capex"]
        )

        diff = (
            calc - hist["fcf"]
        ).abs().max()

        record(
            "HISTORICAL",
            "FCF = CFO - CapEx",
            "PASS" if diff < 0.01 else "FAIL",
            f"Max diff = {diff:.6f}"
        )

    for num, margin in [
        ("ebitda", "ebitda_margin_pct"),
        ("ebit", "ebit_margin_pct"),
        ("net_income", "net_margin_pct"),
        ("fcf", "fcf_margin_pct")
    ]:

        if {
            num,
            margin,
            "revenue"
        }.issubset(hist.columns):

            calc = (
                hist[num]
                / hist["revenue"]
                * 100
            )

            diff = (
                calc - hist[margin]
            ).abs().max()

            record(
                "HISTORICAL",
                f"{margin} reconciliation",
                "PASS" if diff < 0.02 else "FAIL",
                f"Max diff = {diff:.6f} pp"
            )

# ============================================================
# 4. FORECAST
# ============================================================

if forecast is not None:

    record(
        "FORECAST",
        "15 forecast observations",
        "PASS" if len(forecast) == 15 else "FAIL",
        f"Rows = {len(forecast)}"
    )

    if "scenario" in forecast.columns:

        counts = (
            forecast["scenario"]
            .astype(str)
            .str.lower()
            .value_counts()
            .to_dict()
        )

        expected = {
            "bear": 5,
            "base": 5,
            "bull": 5
        }

        record(
            "FORECAST",
            "Five years each Bear/Base/Bull",
            "PASS" if counts == expected else "FAIL",
            counts
        )

    if {
        "revenue",
        "ebit_margin_pct",
        "ebit"
    }.issubset(forecast.columns):

        calc = (
            forecast["revenue"]
            * forecast["ebit_margin_pct"]
            / 100
        )

        diff = (
            calc - forecast["ebit"]
        ).abs().max()

        record(
            "FORECAST",
            "EBIT = Revenue x EBIT margin",
            "PASS" if diff < 1 else "FAIL",
            f"Max diff = {diff:.4f}"
        )

    if {
        "ebit",
        "tax_rate_pct",
        "nopat"
    }.issubset(forecast.columns):

        calc = (
            forecast["ebit"]
            * (
                1 -
                forecast["tax_rate_pct"] / 100
            )
        )

        diff = (
            calc - forecast["nopat"]
        ).abs().max()

        record(
            "FORECAST",
            "NOPAT formula",
            "PASS" if diff < 1 else "FAIL",
            f"Max diff = {diff:.4f}"
        )

    if {
        "nopat",
        "da",
        "capex",
        "delta_nwc",
        "fcff"
    }.issubset(forecast.columns):

        calc = (
            forecast["nopat"]
            + forecast["da"]
            - forecast["capex"]
            - forecast["delta_nwc"]
        )

        diff = (
            calc - forecast["fcff"]
        ).abs().max()

        record(
            "FORECAST",
            "FCFF formula",
            "PASS" if diff < 1 else "FAIL",
            f"Max diff = {diff:.4f}"
        )

# ============================================================
# 5. DCF
# ============================================================

if dcf_final is not None:

    record(
        "DCF",
        "Three DCF scenarios",
        "PASS" if len(dcf_final) == 3 else "FAIL",
        f"Rows = {len(dcf_final)}"
    )

    scenario_col = find_col(
        dcf_final,
        ["scenario"]
    )

    per_share_candidates = [
        c for c in dcf_final.columns
        if "share" in c.lower()
    ]

    if scenario_col and per_share_candidates:

        # Prefer valuation/share column, not shares outstanding.
        preferred = [
            c for c in per_share_candidates
            if (
                "value" in c.lower()
                or "per_share" in c.lower()
                or "price" in c.lower()
            )
            and "outstanding" not in c.lower()
        ]

        share_col = (
            preferred[0]
            if preferred
            else per_share_candidates[-1]
        )

        temp = dcf_final.copy()

        temp[scenario_col] = (
            temp[scenario_col]
            .astype(str)
            .str.lower()
        )

        values = {}

        for s in ["bear", "base", "bull"]:

            row = temp[
                temp[scenario_col] == s
            ]

            if len(row) == 1:
                values[s] = float(
                    row.iloc[0][share_col]
                )

        if len(values) == 3:

            good = (
                values["bear"]
                < values["base"]
                < values["bull"]
            )

            record(
                "DCF",
                "Bear < Base < Bull valuation",
                "PASS" if good else "FAIL",
                values
            )

    # Equity bridge
    ev_col = find_col(
        dcf_final,
        ["enterprise", "value"]
    )

    equity_col = find_col(
        dcf_final,
        ["parent", "equity"]
    )

    net_asset_col = find_col(
        dcf_final,
        ["net", "financial"]
    )

    if (
        ev_col
        and equity_col
        and net_asset_col
    ):

        calc = (
            dcf_final[ev_col]
            + dcf_final[net_asset_col]
        )

        diff = (
            calc - dcf_final[equity_col]
        ).abs().max()

        record(
            "DCF",
            "EV + net financial assets = parent equity",
            "PASS" if diff < 1 else "FAIL",
            f"Max diff = {diff:.4f}"
        )

# ============================================================
# 6. DCF SENSITIVITY
# ============================================================

if sens is not None:

    record(
        "DCF",
        "5x5 sensitivity grid",
        "PASS" if len(sens) == 25 else "FAIL",
        f"Rows = {len(sens)}"
    )

    wacc_col = find_col(
        sens,
        ["wacc"]
    )

    growth_col = find_col(
        sens,
        ["growth"]
    )

    value_candidates = [
        c for c in sens.columns
        if (
            "value" in c.lower()
            or "enterprise" in c.lower()
        )
        and c not in [wacc_col, growth_col]
    ]

    if (
        wacc_col
        and growth_col
        and value_candidates
    ):

        value_col = value_candidates[-1]

        monotonic_ok = True

        for g, group in sens.groupby(growth_col):

            g2 = group.sort_values(wacc_col)

            vals = pd.to_numeric(
                g2[value_col],
                errors="coerce"
            )

            if not vals.is_monotonic_decreasing:
                monotonic_ok = False

        record(
            "DCF",
            "Value falls as WACC rises",
            "PASS" if monotonic_ok else "FAIL"
        )

# ============================================================
# 7. RELATIVE VALUATION
# ============================================================

if comps is not None:

    record(
        "COMPS",
        "Three-company peer table",
        "PASS" if len(comps) == 3 else "FAIL",
        f"Rows = {len(comps)}"
    )

    pe_candidates = [
        c for c in comps.columns
        if (
            "pe" in c.lower()
            or "p/e" in c.lower()
        )
    ]

    if pe_candidates:

        pe_col = pe_candidates[0]

        vals = pd.to_numeric(
            comps[pe_col],
            errors="coerce"
        )

        record(
            "COMPS",
            "Positive P/E values",
            "PASS"
            if vals.notna().all() and (vals > 0).all()
            else "FAIL",
            vals.tolist()
        )

# ============================================================
# 8. PORTFOLIO
# ============================================================

if allocation is not None:

    strategy_col = find_col(
        allocation,
        ["strategy"]
    )

    weight_col = find_col(
        allocation,
        ["weight"]
    )

    money_candidates = [
        c for c in allocation.columns
        if (
            "allocation" in c.lower()
            or "amount" in c.lower()
        )
        and "weight" not in c.lower()
    ]

    if strategy_col and weight_col:

        totals = (
            allocation
            .groupby(strategy_col)[weight_col]
            .sum()
        )

        norm = totals.copy()

        if norm.abs().max() > 2:
            norm = norm / 100

        diff = (
            norm - 1
        ).abs().max()

        record(
            "PORTFOLIO",
            "Strategy weights total 100%",
            "PASS" if diff < 0.001 else "FAIL",
            totals.to_dict()
        )

    if (
        strategy_col
        and money_candidates
    ):

        amount_col = money_candidates[0]

        totals = (
            allocation
            .groupby(strategy_col)[amount_col]
            .sum()
        )

        diff = (
            totals - 1_000_000
        ).abs().max()

        record(
            "PORTFOLIO",
            "Each strategy allocates ₹10 lakh",
            "PASS" if diff < 5 else "FAIL",
            totals.to_dict()
        )

# ============================================================
# 9. PORTFOLIO RISK LOGIC
# ============================================================

if risk is not None:

    var_cols = [
        c for c in risk.columns
        if "var" in c.lower()
        and "cvar" not in c.lower()
    ]

    cvar_cols = [
        c for c in risk.columns
        if "cvar" in c.lower()
    ]

    if var_cols and cvar_cols:

        var = pd.to_numeric(
            risk[var_cols[0]],
            errors="coerce"
        )

        cvar = pd.to_numeric(
            risk[cvar_cols[0]],
            errors="coerce"
        )

        # Both are negative returns in current model.
        good = (cvar <= var).all()

        record(
            "RISK",
            "CVaR is worse than/equal to VaR",
            "PASS" if good else "FAIL"
        )

# ============================================================
# 10. CREDIT SCORECARD
# ============================================================

if credit is not None:

    weight_cols = [
        c for c in credit.columns
        if (
            "weight" in c.lower()
            and "weighted" not in c.lower()
        )
    ]

    if weight_cols:

        total = pd.to_numeric(
            credit[weight_cols[0]],
            errors="coerce"
        ).sum()

        normalized = (
            total / 100
            if total > 2
            else total
        )

        record(
            "CREDIT",
            "Scorecard weights total 100%",
            "PASS"
            if abs(normalized - 1) < 0.001
            else "FAIL",
            f"Raw weight total = {total}"
        )

# ============================================================
# 11. EXCEL INTEGRITY
# ============================================================

xlsx = (
    EXCEL /
    "FinSight_360_TCS_Financial_Model_V2.xlsx"
)

if not xlsx.exists():

    record(
        "EXCEL",
        "V2 workbook exists",
        "FAIL",
        str(xlsx)
    )

else:

    try:
        from openpyxl import load_workbook

        wb = load_workbook(
            xlsx,
            data_only=False
        )

        expected_sheets = {
            "Cover",
            "Assumptions",
            "Valuation Summary",
            "Historical",
            "Forecast",
            "DCF",
            "Sensitivity",
            "Comps",
            "FP&A",
            "Treasury",
            "Credit",
            "Portfolio",
            "Risk"
        }

        missing = (
            expected_sheets
            - set(wb.sheetnames)
        )

        record(
            "EXCEL",
            "13 required worksheets",
            "PASS" if not missing else "FAIL",
            (
                "All present"
                if not missing
                else sorted(missing)
            )
        )

        formula_count = 0
        bad_refs = []

        for ws in wb.worksheets:

            for row in ws.iter_rows():

                for cell in row:

                    value = cell.value

                    if (
                        isinstance(value, str)
                        and value.startswith("=")
                    ):
                        formula_count += 1

                    if isinstance(value, str):

                        if any(
                            x in value
                            for x in [
                                "#REF!",
                                "#DIV/0!",
                                "#VALUE!",
                                "#NAME?"
                            ]
                        ):
                            bad_refs.append(
                                f"{ws.title}!{cell.coordinate}"
                            )

        record(
            "EXCEL",
            "Workbook contains formulas",
            "PASS" if formula_count > 0 else "FAIL",
            f"Formula count = {formula_count}"
        )

        record(
            "EXCEL",
            "No obvious broken formula references",
            "PASS" if not bad_refs else "FAIL",
            (
                "None"
                if not bad_refs
                else bad_refs
            )
        )

    except Exception as e:

        record(
            "EXCEL",
            "Workbook integrity",
            "FAIL",
            str(e)
        )

# ============================================================
# 12. POWER BI DATA MART
# ============================================================

required_pbi = [
    "financial_performance.csv",
    "financial_forecast.csv",
    "peer_valuation.csv",
    "dcf_valuation.csv",
    "dcf_sensitivity.csv",
    "portfolio_strategies.csv",
    "portfolio_allocation.csv",
    "portfolio_risk.csv",
    "portfolio_stress.csv",
    "fpa_variance.csv",
    "treasury_kpis.csv",
    "treasury_stress.csv",
    "credit_scorecard.csv",
    "credit_summary.csv"
]

for name in required_pbi:

    path = PBI / name

    if not path.exists():

        record(
            "POWER BI",
            name,
            "FAIL",
            "Missing"
        )

    else:

        try:
            df = pd.read_csv(path)

            record(
                "POWER BI",
                name,
                "PASS" if len(df) > 0 else "FAIL",
                f"{len(df)} rows"
            )

        except Exception as e:

            record(
                "POWER BI",
                name,
                "FAIL",
                str(e)
            )

# ============================================================
# 13. KNOWN DATA / METHODOLOGY GAPS
# ============================================================

known_warnings = [
    (
        "HCLTECH FY2024",
        "Some standardized FY2024 fields remain unavailable. "
        "Recovered CFF -15464 still requires final DB/CSV reconciliation."
    ),
    (
        "HCLTECH Balance Sheet",
        "Some detailed receivables/PPE/payables fields are not yet standardized."
    ),
    (
        "Net Income Attribution",
        "Consolidated PAT vs attributable-to-parent treatment needs final consistency review."
    ),
    (
        "Equity Attribution",
        "Shareholders' equity vs total equity treatment needs final consistency review."
    ),
    (
        "Sharpe / Sortino",
        "Current market and portfolio modules use a 0% risk-free baseline."
    ),
    (
        "DCF WACC",
        "12.25% is currently a modelling assumption/equity-dominant proxy; final methodology/source note required."
    ),
    (
        "Terminal Growth",
        "Terminal growth rates are modelling assumptions, not observed company guidance."
    ),
    (
        "DCF Net Financial Assets",
        "Treatment of investments/deposits as excess financial assets must be explicitly documented."
    ),
    (
        "Lease Treatment",
        "Lease liabilities are deducted in the equity bridge while FCFF is not explicitly lease-adjusted."
    ),
    (
        "DCF Timing",
        "DCF anchor is FY2026 year-end while market reference price is 06-Oct-2026."
    ),
    (
        "FCF vs FCFF",
        "Historical FCF and forecast FCFF are different concepts and must not be presented as identical."
    ),
    (
        "Relative Valuation",
        "EV/EBITDA is intentionally omitted until standardized peer net-debt/share inputs are available."
    ),
    (
        "Credit Score",
        "Internal 98% credit score is analytical only and is not an external agency rating."
    ),
    (
        "Portfolio Optimization",
        "Maximum-Sharpe portfolio is highly concentrated in HCLTECH; concentration risk must be disclosed."
    )
]

for test, detail in known_warnings:

    record(
        "METHODOLOGY",
        test,
        "WARNING",
        detail
    )

# ============================================================
# 14. SAVE REPORT
# ============================================================

report = pd.DataFrame(audit)

output = R / "master_project_audit.csv"

report.to_csv(
    output,
    index=False
)

counts = (
    report["status"]
    .value_counts()
    .to_dict()
)

passes = counts.get("PASS", 0)
warnings = counts.get("WARNING", 0)
fails = counts.get("FAIL", 0)

print("\nAUDIT SUMMARY")
print("-" * 100)

print(f"PASS    : {passes}")
print(f"WARNING : {warnings}")
print(f"FAIL    : {fails}")

print("\nFAILURES")
print("-" * 100)

failure_rows = report[
    report["status"] == "FAIL"
]

if failure_rows.empty:
    print("No hard failures detected.")
else:
    for _, row in failure_rows.iterrows():
        print(
            f"[FAIL] {row['area']} | "
            f"{row['test']} | "
            f"{row['detail']}"
        )

print("\nWARNINGS")
print("-" * 100)

warning_rows = report[
    report["status"] == "WARNING"
]

if warning_rows.empty:
    print("No warnings.")
else:
    for _, row in warning_rows.iterrows():
        print(
            f"[WARNING] {row['area']} | "
            f"{row['test']} | "
            f"{row['detail']}"
        )

print("\n" + "=" * 100)
print(f"Audit report saved -> {output}")
print("=" * 100)

if fails == 0:

    print(
        "\nRESULT: MATHEMATICAL/STRUCTURAL AUDIT PASSED "
        "WITH METHODOLOGY WARNINGS."
    )

    print(
        "NEXT: Resolve warnings before final Power BI/reporting."
    )

else:

    print(
        "\nRESULT: HARD FAILURES FOUND."
    )

    print(
        "DO NOT PROCEED. Fix failures first."
    )
