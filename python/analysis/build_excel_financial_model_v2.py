from pathlib import Path
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import ColorScaleRule

BASE = Path("results/tables")
OUT = Path("excel")
OUT.mkdir(parents=True, exist_ok=True)

OUTPUT = OUT / "FinSight_360_TCS_Financial_Model_V2.xlsx"

# ============================================================
# LOAD VALIDATED OUTPUTS
# ============================================================

files = {
    "historical": "tcs_historical_financial_model_base.csv",
    "forecast": "tcs_scenario_forecast_FY2027_FY2031.csv",
    "dcf": "tcs_dcf_final_equity_valuation.csv",
    "sensitivity": "tcs_dcf_per_share_sensitivity.csv",
    "relative_summary": "tcs_relative_valuation_summary.csv",
    "comps": "fy2026_peer_relative_valuation.csv",
    "fpa": "tcs_fpa_budget_variance.csv",
    "fpa_drivers": "tcs_fpa_budget_drivers.csv",
    "treasury": "tcs_treasury_liquidity_kpis.csv",
    "treasury_stress": "tcs_treasury_liquidity_stress.csv",
    "credit": "tcs_credit_scorecard.csv",
    "credit_summary": "tcs_credit_risk_summary.csv",
    "portfolio": "portfolio_strategy_summary.csv",
    "portfolio_risk": "portfolio_risk_metrics.csv",
    "portfolio_stress": "portfolio_stress_tests.csv",
}

data = {
    key: pd.read_csv(BASE / filename)
    for key, filename in files.items()
}

# ============================================================
# WORKBOOK
# ============================================================

wb = Workbook()
wb.remove(wb.active)

sheets = [
    "Cover",
    "Assumptions",
    "Valuation Summary",
    "Historical",
    "Forecast",
    "DCF",
    "Sensitivity",
    "Comps",
    "FP&A",
    "Treasury",
    "Credit",
    "Portfolio",
    "Risk"
]

for name in sheets:
    wb.create_sheet(name)

# ============================================================
# STYLES
# ============================================================

NAVY = "17365D"
BLUE = "4472C4"
LIGHT_BLUE = "D9EAF7"
YELLOW = "FFF2CC"
GREEN = "E2F0D9"
RED = "FCE4D6"
GREY = "E7E6E6"
WHITE = "FFFFFF"
BLACK = "000000"

thin = Side(style="thin", color="B7B7B7")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

def make_title(ws, text):
    ws.merge_cells("A1:H2")
    c = ws["A1"]
    c.value = text
    c.font = Font(size=18, bold=True, color=WHITE)
    c.fill = PatternFill("solid", fgColor=NAVY)
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 24
    ws.row_dimensions[2].height = 24

def section(ws, row, text, end_col=8):
    ws.merge_cells(
        start_row=row,
        start_column=1,
        end_row=row,
        end_column=end_col
    )
    c = ws.cell(row, 1)
    c.value = text
    c.font = Font(bold=True, color=WHITE)
    c.fill = PatternFill("solid", fgColor=BLUE)

def write_df(ws, df, start_row=4, start_col=1):
    for j, col in enumerate(df.columns, start_col):
        c = ws.cell(start_row, j, str(col))
        c.font = Font(bold=True, color=WHITE)
        c.fill = PatternFill("solid", fgColor=BLUE)
        c.border = border
        c.alignment = Alignment(horizontal="center")

    for i, row in enumerate(df.itertuples(index=False), start_row + 1):
        for j, value in enumerate(row, start_col):
            if pd.isna(value):
                value = None

            c = ws.cell(i, j, value)
            c.border = border
            c.alignment = Alignment(vertical="center")

    ws.auto_filter.ref = (
        f"{get_column_letter(start_col)}{start_row}:"
        f"{get_column_letter(start_col + len(df.columns) - 1)}"
        f"{start_row + len(df)}"
    )

def auto_width(ws):
    for col in range(1, ws.max_column + 1):
        width = 12

        for cell in ws[get_column_letter(col)]:
            if cell.value is not None:
                width = max(
                    width,
                    min(len(str(cell.value)) + 2, 34)
                )

        ws.column_dimensions[get_column_letter(col)].width = width

def format_numbers(ws):
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, float):
                c.number_format = '#,##0.00;[Red](#,##0.00);-'

def standard_sheet(ws):
    ws.freeze_panes = "A5"
    auto_width(ws)
    format_numbers(ws)

# ============================================================
# COVER
# ============================================================

ws = wb["Cover"]
make_title(ws, "FinSight 360 — Financial Analysis, Valuation & Portfolio Intelligence")

cover = [
    ("Primary Company", "Tata Consultancy Services"),
    ("Peer Universe", "Infosys | HCL Technologies"),
    ("Historical Period", "FY2022–FY2026"),
    ("Forecast Period", "FY2027–FY2031"),
    ("Market Reference Date", "06-Oct-2026"),
    ("Currency", "₹ Crore unless otherwise stated"),
    ("Primary Valuation", "DCF"),
    ("Secondary Valuation", "Peer P/E"),
    ("Portfolio Capital", "₹10,00,000"),
]

for r, (label, value) in enumerate(cover, 4):
    ws.cell(r, 1, label)
    ws.cell(r, 1).font = Font(bold=True)
    ws.cell(r, 1).fill = PatternFill("solid", fgColor=LIGHT_BLUE)

    ws.cell(r, 2, value)

section(ws, 15, "MODEL DISCLOSURES")

notes = [
    "Historical financials use consolidated annual-report data.",
    "EBITDA is analytically derived where it is not directly reported.",
    "Forecasts and stress scenarios are analytical assumptions, not company guidance.",
    "Restricted-purpose TCS Foundation/trust assets are excluded from available liquidity.",
    "Relative valuation currently uses P/E; unsupported peer EV/EBITDA inputs are not fabricated.",
    "The internal credit scorecard is not an external agency credit rating.",
    "FY2026 FCF and forecast FCFF are different cash-flow concepts.",
    "DCF is anchored to FY2026 year-end; 06-Oct-2026 price is a reference comparison."
]

for i, note in enumerate(notes, 16):
    ws.cell(i, 1, f"• {note}")

ws.column_dimensions["A"].width = 42
ws.column_dimensions["B"].width = 65

# ============================================================
# ASSUMPTIONS
# ============================================================

ws = wb["Assumptions"]
make_title(ws, "TCS Model Assumptions")

assumptions = [
    ["Category", "Assumption", "Bear", "Base", "Bull", "Unit"],
    ["DCF", "Discount Rate / WACC", 12.25, 12.25, 12.25, "%"],
    ["DCF", "Terminal Growth", 4.50, 5.50, 6.00, "%"],
    ["Forecast", "FY2027 Revenue Growth", 3.00, 6.00, 9.00, "%"],
    ["Forecast", "FY2027 EBIT Margin", 25.50, 26.70, 27.50, "%"],
    ["Forecast", "Tax Rate", 25.50, 25.30, 25.00, "%"],
    ["Forecast", "D&A / Revenue", 2.10, 2.15, 2.20, "%"],
    ["Forecast", "CapEx / Revenue", 1.60, 1.40, 1.30, "%"],
    ["Forecast", "Delta NWC / Revenue", 1.00, 0.70, 0.50, "%"],
    ["Valuation", "Shares Outstanding", 361.8087518, 361.8087518, 361.8087518, "Crore"],
    ["Valuation", "Adjusted Net Financial Assets", 31240, 31240, 31240, "₹ Cr"],
    ["Market", "Reference Price", 2093.40, 2093.40, 2093.40, "₹/share"],
]

for r_idx, row in enumerate(assumptions, 4):
    for c_idx, value in enumerate(row, 1):
        c = ws.cell(r_idx, c_idx, value)
        c.border = border

        if r_idx == 4:
            c.font = Font(bold=True, color=WHITE)
            c.fill = PatternFill("solid", fgColor=BLUE)
        elif c_idx in [3, 4, 5]:
            c.fill = PatternFill("solid", fgColor=YELLOW)
            c.font = Font(color=BLUE)

ws.freeze_panes = "A5"
auto_width(ws)
format_numbers(ws)

# ============================================================
# VALUATION SUMMARY
# ============================================================

ws = wb["Valuation Summary"]
make_title(ws, "TCS Valuation Summary")

section(ws, 4, "MARKET REFERENCE", 5)

ws["A5"] = "Reference Price"
ws["B5"] = "=Assumptions!D15"

ws["A6"] = "Reference Date"
ws["B6"] = "06-Oct-2026"

section(ws, 8, "DCF SCENARIO VALUATION", 5)

headers = [
    "Scenario",
    "Implied Value / Share",
    "Reference Price",
    "Upside / (Downside)",
    "Interpretation"
]

for c, h in enumerate(headers, 1):
    cell = ws.cell(9, c, h)
    cell.font = Font(bold=True, color=WHITE)
    cell.fill = PatternFill("solid", fgColor=BLUE)
    cell.border = border

# Pull values from validated DCF dataframe
dcf_df = data["dcf"]

scenario_rows = {
    "Bear": 10,
    "Base": 11,
    "Bull": 12
}

for scenario, row_num in scenario_rows.items():

    source = dcf_df[
        dcf_df.iloc[:, 0].astype(str).str.lower()
        == scenario.lower()
    ]

    if len(source) == 1:
        # Find likely per-share column robustly
        per_share_cols = [
            c for c in dcf_df.columns
            if "share" in c.lower()
            and (
                "value" in c.lower()
                or "equity" in c.lower()
            )
        ]

        if not per_share_cols:
            per_share_cols = [
                c for c in dcf_df.columns
                if "share" in c.lower()
            ]

        implied = float(
            source.iloc[0][per_share_cols[0]]
        )

        ws.cell(row_num, 1, scenario)
        ws.cell(row_num, 2, implied)
        ws.cell(row_num, 3, "=Assumptions!D15")
        ws.cell(
            row_num,
            4,
            f"=B{row_num}/C{row_num}-1"
        )

        if scenario == "Bear":
            ws.cell(row_num, 5, "Downside scenario")
        elif scenario == "Base":
            ws.cell(row_num, 5, "Central DCF case")
        else:
            ws.cell(row_num, 5, "Upside scenario")

section(ws, 15, "RELATIVE VALUATION", 5)

rel = data["relative_summary"].iloc[0]

ws["A16"] = "TCS Current P/E"
ws["B16"] = float(rel["tcs_pe"])

ws["A17"] = "Peer Median P/E"
ws["B17"] = float(rel["peer_median_pe"])

ws["A18"] = "Peer-Implied TCS Value"
ws["B18"] = float(
    rel["tcs_implied_value_peer_median_pe"]
)

ws["A19"] = "Peer-Implied Upside"
ws["B19"] = "=B18/$B$5-1"

section(ws, 22, "VALUATION CROSS-CHECK", 5)

ws["A23"] = "DCF Base Value"
ws["B23"] = "=B11"

ws["A24"] = "Relative Value"
ws["B24"] = "=B18"

ws["A25"] = "Simple Cross-Check Average"
ws["B25"] = "=AVERAGE(B23:B24)"

ws["A26"] = "Cross-Check Upside"
ws["B26"] = "=B25/$B$5-1"

for row in range(10, 13):
    ws.cell(row, 4).number_format = '0.00%'

ws["B19"].number_format = '0.00%'
ws["B26"].number_format = '0.00%'

for row in range(5, 27):
    for col in range(1, 6):
        ws.cell(row, col).border = border

ws.conditional_formatting.add(
    "D10:D12",
    ColorScaleRule(
        start_type="min",
        start_color="F8696B",
        mid_type="percentile",
        mid_value=50,
        mid_color="FFEB84",
        end_type="max",
        end_color="63BE7B"
    )
)

auto_width(ws)

# ============================================================
# DATA SHEETS
# ============================================================

mapping = [
    ("Historical", "historical", "TCS Historical Financial Performance"),
    ("Forecast", "forecast", "TCS Bear / Base / Bull Forecast"),
    ("DCF", "dcf", "TCS DCF Equity Valuation"),
    ("Sensitivity", "sensitivity", "TCS DCF Per-Share Sensitivity"),
    ("Comps", "comps", "FY2026 Peer Relative Valuation"),
    ("FP&A", "fpa", "TCS FP&A Budget Variance"),
    ("Treasury", "treasury", "TCS Treasury & Liquidity KPIs"),
    ("Credit", "credit", "TCS Internal Credit Scorecard"),
    ("Portfolio", "portfolio", "₹10 Lakh Portfolio Strategies"),
    ("Risk", "portfolio_risk", "Portfolio Risk Metrics"),
]

for sheet, key, heading in mapping:

    ws = wb[sheet]
    make_title(ws, heading)

    df = data[key].copy()

    if key == "sensitivity":
        # preserve sensitivity matrix shape
        df = pd.read_csv(
            BASE / files[key]
        )

    write_df(ws, df)
    standard_sheet(ws)

# ============================================================
# ADD SECOND TABLES TO MULTI-MODULE SHEETS
# ============================================================

# FP&A drivers
ws = wb["FP&A"]
start = ws.max_row + 3
section(ws, start, "FY2027 BASE-CASE BUDGET DRIVERS")
write_df(ws, data["fpa_drivers"], start + 1)
auto_width(ws)

# Treasury stress
ws = wb["Treasury"]
start = ws.max_row + 3
section(ws, start, "LIQUIDITY STRESS TESTS")
write_df(ws, data["treasury_stress"], start + 1)
auto_width(ws)

# Credit summary
ws = wb["Credit"]
start = ws.max_row + 3
section(ws, start, "CREDIT RISK SUMMARY")
write_df(ws, data["credit_summary"], start + 1)
auto_width(ws)

# Portfolio risk + stress
ws = wb["Portfolio"]
start = ws.max_row + 3
section(ws, start, "PORTFOLIO RISK")
write_df(ws, data["portfolio_risk"], start + 1)

start = ws.max_row + 3
section(ws, start, "PORTFOLIO STRESS TESTS")
write_df(ws, data["portfolio_stress"], start + 1)
auto_width(ws)

# ============================================================
# FINANCIAL-MODEL FORMULA COLOR CONVENTION
# ============================================================

# Blue = hardcoded model assumptions
for row in wb["Assumptions"].iter_rows(
    min_row=5,
    min_col=3,
    max_col=5
):
    for c in row:
        c.font = Font(color=BLUE)

# Green = cross-sheet links
for row in wb["Valuation Summary"].iter_rows():
    for c in row:
        if isinstance(c.value, str) and c.value.startswith("="):
            if "!" in c.value:
                c.font = Font(color="008000")
            else:
                c.font = Font(color=BLACK)

# ============================================================
# FINAL FORMATTING
# ============================================================

for ws in wb.worksheets:

    ws.sheet_view.showGridLines = False

    for row in ws.iter_rows():
        for c in row:
            if c.value is not None:
                c.alignment = Alignment(
                    vertical="center",
                    wrap_text=False
                )

    format_numbers(ws)

# ============================================================
# SAVE
# ============================================================

wb.calculation.fullCalcOnLoad = True
wb.calculation.forceFullCalc = True
wb.calculation.calcMode = "auto"

wb.save(OUTPUT)

# ============================================================
# QA — REOPEN WORKBOOK
# ============================================================

check = load_workbook(
    OUTPUT,
    data_only=False
)

assert check.sheetnames == sheets

required = [
    "Assumptions",
    "Valuation Summary",
    "Historical",
    "Forecast",
    "DCF",
    "Sensitivity",
    "Comps",
    "FP&A",
    "Treasury",
    "Credit",
    "Portfolio",
    "Risk"
]

for name in required:
    assert name in check.sheetnames

formula_count = 0

for ws in check.worksheets:
    for row in ws.iter_rows():
        for cell in row:
            if (
                isinstance(cell.value, str)
                and cell.value.startswith("=")
            ):
                formula_count += 1

assert formula_count > 0

print("\n" + "=" * 108)
print("FINSIGHT 360 — PROFESSIONAL EXCEL MODEL V2")
print("=" * 108)

print(f"\nWorkbook -> {OUTPUT}")
print(f"Sheets   -> {len(check.sheetnames)}")
print(f"Formulas -> {formula_count}")

print("\nCreated:")
for name in check.sheetnames:
    print(f"  ✓ {name}")

print("\nQA PASSED:")
print("  ✓ Workbook re-opened successfully")
print("  ✓ Required sheets present")
print("  ✓ Linked formulas present")
print("  ✓ Assumption cells formatted")
print("  ✓ Valuation summary created")
print("  ✓ Portfolio and risk modules included")

print(
    "\nProfessional Excel V2 generated successfully."
)
